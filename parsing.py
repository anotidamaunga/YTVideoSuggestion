import io

from pptx import Presentation
from pypdf import PdfReader
from docx import Document

supported_extensions = {"pdf", "pptx", "docx", "md"}

def parse_documents(filename: str, file_bytes: bytes) -> str:
    ext = filename.split(".", 1)[-1].lower()
    if ext not in supported_extensions:
        raise ValueError("Unsupported file extension")

    if ext == "pdf":
        return parse_pdf(file_bytes)
    if ext == "pptx":
        return parse_pptx(file_bytes)
    if ext == "docx":
        return parse_docx(file_bytes)

    return file_bytes.decode("utf-8")

def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return  "\n\n".join(pages)

def parse_pptx(file_bytes: bytes) -> str:
    pptx_reader = Presentation(io.BytesIO(file_bytes))
    slides = []

    for slide in pptx_reader.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        if texts:
            slides.append("\n".join(texts))

    return "\n".join(slides)



def parse_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            prefix = "# " if para.style.name.lower().startswith("heading") else ""
            parts.append(prefix + para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)